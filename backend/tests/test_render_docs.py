from docx import Document
from pypdf import PdfReader

from scripts.render_docs import (
    SourceDoc,
    read_source,
    render_docx,
    render_html,
    render_pdf,
    to_blocks,
)

SAMPLE = """# Widget Guide

Intro with **bold** and *italic* text.

## Steps

1. Plug it in
2. Press the button

- Fast
- Quiet

| Spec | Value |
| --- | --- |
| Battery | 10 hours |
"""


def _doc(tmp_path) -> SourceDoc:
    src = tmp_path / "widget-guide.md"
    front_matter = "---\ntitle: Widget Guide\ndoc_type: manual\nupdated: 2026-09-01\n---\n"
    src.write_text(front_matter + SAMPLE)
    return read_source(src)


def test_front_matter_is_parsed(tmp_path):
    doc = _doc(tmp_path)

    assert doc.meta == {"title": "Widget Guide", "doc_type": "manual", "updated": "2026-09-01"}
    assert doc.body.startswith("# Widget Guide")


def test_markdown_becomes_blocks():
    blocks = to_blocks(SAMPLE)

    assert [b.kind for b in blocks] == ["heading", "paragraph", "heading", "list", "list", "table"]
    bold = [r for r in blocks[1].runs if r.bold]
    assert bold[0].text == "bold"
    assert blocks[3].ordered and not blocks[4].ordered
    assert len(blocks[5].rows) == 2


def test_pdf_contains_text_and_title(tmp_path):
    out = tmp_path / "w.pdf"
    render_pdf(_doc(tmp_path), out)

    reader = PdfReader(out)
    text = "".join(page.extract_text() for page in reader.pages)
    assert reader.metadata.title == "Widget Guide"
    assert "Press the button" in text and "10 hours" in text


def test_docx_contains_headings_lists_and_table(tmp_path):
    out = tmp_path / "w.docx"
    render_docx(_doc(tmp_path), out)

    document = Document(out)
    styles = {p.style.name for p in document.paragraphs}
    assert document.core_properties.title == "Widget Guide"
    assert {"Title", "Heading 1", "List Number", "List Bullet"} <= styles
    assert document.tables[0].cell(1, 1).text == "10 hours"


def test_html_wraps_content_in_page_template(tmp_path):
    out = tmp_path / "w.html"
    render_html(_doc(tmp_path), out)

    page = out.read_text()
    assert "<title>Widget Guide | VoltCart Help Center</title>" in page
    assert "<table>" in page and "<strong>bold</strong>" in page
    assert "<footer>" in page
