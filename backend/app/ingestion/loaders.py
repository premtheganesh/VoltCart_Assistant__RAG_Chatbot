"""Parse PDF, DOCX and HTML files into one common structure: an ordered list of blocks.

Every loader keeps the document's structure (headings, lists, tables) because chunking and
retrieval quality depend on it: a table flattened into one line, or a heading lost, makes the
matching chunk much harder to find.
"""

import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Literal

import pdfplumber
from bs4 import BeautifulSoup, Tag
from docx import Document as DocxDocument
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph


@dataclass
class Block:
    kind: Literal["heading", "text", "table"]
    text: str
    level: int = 0  # headings only: 1 = document title, 2 = section, 3 = subsection
    page: int | None = None


@dataclass
class ParsedDocument:
    source: str
    title: str
    doc_type: str
    blocks: list[Block]
    url: str | None = None
    updated: date | None = None
    products: list[str] = field(default_factory=list)


def normalize_space(text: str) -> str:
    return re.sub(r"\s+([.,;:!?)])", r"\1", re.sub(r"\s+", " ", text)).strip()


def table_to_markdown(rows: list[list[str]]) -> str:
    """Rows of cell strings (first row = header) to a Markdown table."""
    cleaned = [[normalize_space(c or "").replace("|", "/") for c in row] for row in rows if row]
    width = max(len(r) for r in cleaned)
    cleaned = [r + [""] * (width - len(r)) for r in cleaned]
    lines = ["| " + " | ".join(cleaned[0]) + " |", "|" + " --- |" * width]
    lines += ["| " + " | ".join(r) + " |" for r in cleaned[1:]]
    return "\n".join(lines)


def _split_keywords(value: str | None) -> list[str]:
    return [k.strip() for k in (value or "").split(",") if k.strip()]


def _parse_date(value: str | None) -> date | None:
    try:
        return date.fromisoformat((value or "")[:10])
    except ValueError:
        return None


# ---------- PDF ----------

BULLET = re.compile(r"^(\(cid:127\)|[•▪●◦-])\s*")
NUMBER = re.compile(r"^(\d+)[.)]?\s+")
PDF_MARGIN = 0.07  # top/bottom share of the page holding running headers, footers, page numbers


def _numbered_item(line: dict) -> str | None:
    """'3 Bend the hook' with a wide gap after the number is a list item, '16 GB of RAM' is not."""
    match = NUMBER.match(line["text"])
    if not match:
        return None
    chars = line["chars"]
    digits = len(match.group(1))
    gap = chars[digits]["x0"] - chars[digits - 1]["x1"] if len(chars) > digits else 0
    if line["text"][digits] in ".)" or gap > chars[0]["size"] * 0.5:
        return f"{match.group(1)}. {line['text'][match.end() :].strip()}"
    return None


def _inside(line: dict, box: tuple[float, float, float, float]) -> bool:
    x0, top, x1, bottom = box
    return line["top"] >= top - 1 and line["bottom"] <= bottom + 1 and line["x0"] >= x0 - 1


def _page_blocks(lines, tables, page_no, heading_size, level_of) -> list[Block]:
    """One PDF page: merge wrapped lines into paragraphs and list items, keep tables whole."""
    blocks: list[Block] = []
    paragraph: list[str] = []
    list_items: list[str] = []
    prev_bottom = None

    def flush() -> None:
        if paragraph:
            blocks.append(Block("text", " ".join(paragraph), page=page_no))
            paragraph.clear()
        if list_items:
            blocks.append(Block("text", "\n".join(list_items), page=page_no))
            list_items.clear()

    items = [("line", ln["top"], ln) for ln in lines]
    items += [("table", bbox[1], rows) for bbox, rows in tables]
    for kind, _, item in sorted(items, key=lambda i: i[1]):
        if kind == "table":
            flush()
            if item:
                blocks.append(Block("table", table_to_markdown(item), page=page_no))
            prev_bottom = None
            continue
        text = normalize_space(item["text"])
        if not text:
            continue
        if size := heading_size(item):
            flush()
            blocks.append(Block("heading", text, level=level_of[size], page=page_no))
        elif (numbered := _numbered_item(item)) or BULLET.match(text):
            if paragraph:
                flush()
            list_items.append(numbered or BULLET.sub("- ", text))
        elif list_items and prev_bottom is not None and item["top"] - prev_bottom < 4:
            list_items[-1] += " " + text  # wrapped line of a list item
        else:
            gap = item["top"] - prev_bottom if prev_bottom is not None else 99
            if list_items or gap > item["chars"][0]["size"] * 0.8:
                flush()
            paragraph.append(text)
        prev_bottom = item["bottom"]
    flush()
    return blocks


def load_pdf(path: Path, source: str) -> ParsedDocument:
    with pdfplumber.open(path) as pdf:
        meta = pdf.metadata
        pages = []  # per page: (lines outside tables, tables)
        for page in pdf.pages:
            low, high = page.height * PDF_MARGIN, page.height * (1 - PDF_MARGIN)
            tables = [(t.bbox, t.extract()) for t in page.find_tables()]
            lines = [
                ln
                for ln in page.extract_text_lines()
                if ln["top"] > low
                and ln["bottom"] < high
                and not any(_inside(ln, bbox) for bbox, _ in tables)
            ]
            pages.append((lines, tables))

    # Body text = the most common font size; bold lines clearly larger than it are headings
    sizes = Counter(round(c["size"], 1) for lines, _ in pages for ln in lines for c in ln["chars"])
    body_size = sizes.most_common(1)[0][0] if sizes else 10.0

    def heading_size(line: dict) -> float | None:
        first = line["chars"][0]
        if "Bold" in first["fontname"] and first["size"] > body_size * 1.15:
            return round(first["size"], 1)
        return None

    levels = sorted(
        {s for lines, _ in pages for ln in lines if (s := heading_size(ln))}, reverse=True
    )
    level_of = {s: min(i + 1, 3) for i, s in enumerate(levels)}

    blocks: list[Block] = []
    for page_no, (lines, tables) in enumerate(pages, start=1):
        blocks += _page_blocks(lines, tables, page_no, heading_size, level_of)

    title = meta.get("Title") or next((b.text for b in blocks if b.kind == "heading"), path.stem)
    return ParsedDocument(
        source=source,
        title=title,
        doc_type=meta.get("Subject") or "document",
        blocks=blocks,
        products=_split_keywords(meta.get("Keywords")),
    )


# ---------- DOCX ----------


def load_docx(path: Path, source: str) -> ParsedDocument:
    document = DocxDocument(str(path))
    props = document.core_properties
    blocks: list[Block] = []
    list_items: list[str] = []

    def flush_list() -> None:
        if list_items:
            blocks.append(Block("text", "\n".join(list_items)))
            list_items.clear()

    # Walk the body in order, so tables stay between the paragraphs they belong to
    for element in document.element.body.iterchildren():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "tbl":
            flush_list()
            table = DocxTable(element, document)
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            if rows:
                blocks.append(Block("table", table_to_markdown(rows)))
            continue
        if tag != "p":
            continue
        paragraph = DocxParagraph(element, document)
        text = normalize_space(paragraph.text)
        if not text:
            continue
        style = paragraph.style.name if paragraph.style is not None else ""
        if style == "Title":
            flush_list()
            blocks.append(Block("heading", text, level=1))
        elif style.startswith("Heading"):
            flush_list()
            level = int(style.split()[-1]) + 1 if style.split()[-1].isdigit() else 2
            blocks.append(Block("heading", text, level=min(level, 3)))
        elif style.startswith("List"):
            prefix = f"{len(list_items) + 1}. " if "Number" in style else "- "
            list_items.append(prefix + text)
        else:
            flush_list()
            blocks.append(Block("text", text))
    flush_list()

    title = props.title or next((b.text for b in blocks if b.kind == "heading"), path.stem)
    return ParsedDocument(
        source=source,
        title=title,
        doc_type=props.subject or "document",
        blocks=blocks,
        updated=props.modified.date() if props.modified else None,
        products=_split_keywords(props.keywords),
    )


# ---------- HTML ----------

BOILERPLATE = ["script", "style", "noscript", "header", "nav", "footer", "aside", "form"]
# Page chrome inside the main content; the date is kept as metadata instead
CHROME_SELECTORS = ".breadcrumb, [aria-label=breadcrumb], .updated, .last-updated"
CONTENT_TAGS = ["h1", "h2", "h3", "h4", "p", "ul", "ol", "table", "pre"]


def load_html(path: Path, source: str) -> ParsedDocument:
    soup = BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser")

    def meta(name: str) -> str | None:
        tag = soup.find("meta", attrs={"name": name})
        return tag.get("content") if isinstance(tag, Tag) else None

    canonical = soup.find("link", rel="canonical")
    page_title = soup.title.get_text(strip=True).split(" | ")[0] if soup.title else path.stem

    # Remove site chrome (menus, footers, breadcrumbs) so it never pollutes chunks
    for tag in soup.find_all(BOILERPLATE):
        tag.decompose()
    for tag in soup.select(CHROME_SELECTORS):
        tag.decompose()
    root = soup.find("main") or soup.find("article") or soup.body or soup

    blocks: list[Block] = []
    for el in root.find_all(CONTENT_TAGS):
        # Skip elements nested in a list or table: the container is emitted as a whole
        if el.find_parent(["ul", "ol", "table"]) is not None:
            continue
        if el.name in ("h1", "h2", "h3", "h4"):
            text = normalize_space(el.get_text(" "))
            if text:
                blocks.append(Block("heading", text, level=min(int(el.name[1]), 3)))
        elif el.name in ("ul", "ol"):
            items = [normalize_space(li.get_text(" ")) for li in el.find_all("li", recursive=False)]
            numbered = el.name == "ol"
            lines = [f"{i}. {t}" if numbered else f"- {t}" for i, t in enumerate(items, 1) if t]
            if lines:
                blocks.append(Block("text", "\n".join(lines)))
        elif el.name == "table":
            rows = [
                [cell.get_text(" ") for cell in tr.find_all(["th", "td"])]
                for tr in el.find_all("tr")
            ]
            if rows:
                blocks.append(Block("table", table_to_markdown(rows)))
        else:
            text = normalize_space(el.get_text(" "))
            if text:
                blocks.append(Block("text", text))

    h1 = next((b.text for b in blocks if b.kind == "heading" and b.level == 1), None)
    return ParsedDocument(
        source=source,
        title=h1 or page_title,
        doc_type=meta("doc_type") or "web_page",
        blocks=blocks,
        url=canonical.get("href") if isinstance(canonical, Tag) else None,
        updated=_parse_date(meta("updated")),
    )


LOADERS = {".pdf": load_pdf, ".docx": load_docx, ".html": load_html, ".htm": load_html}


def load_file(path: Path, source: str) -> ParsedDocument:
    loader = LOADERS.get(path.suffix.lower())
    if loader is None:
        raise ValueError(f"Unsupported file type: {path.name}")
    return loader(path, source)
